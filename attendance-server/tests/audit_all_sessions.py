"""Read-only sweep: every session's sheet vs the database, across the workbook.

Diagnostic, not part of the running system — run it by hand to see the whole
picture. The live verifier in app.py only checks the ACTIVE sessions, because
reading all 164 tabs on a loop would fight with attendance writes for the file.

Sheet names MUST be sanitized before lookup: Excel caps tab names at 31 chars,
so the session `2026-03-30_Bz 23rd March Evening` (32) lives on the tab
`2026-03-30_Bz 23rd March Evenin`. Looking up the raw session id instead reports
a perfectly healthy session as having no sheet at all.

Usage:  python tests/audit_all_sessions.py
"""
import os
import sqlite3
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, ".."))
sys.path.insert(0, ROOT)

import session_sync as S  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

LIVE = os.path.join(ROOT, "Server file", "live_attendance_master.xlsx")
DB = os.path.join(ROOT, "database", "attendance.db")

_BAD = {"/": "-", "\\": "-", ":": "-", "?": "", "*": ""}


def sanitize_sheet_name(name):
    """Mirror of app.sanitize_sheet_name — kept in step by test_session_sync."""
    text = str(name or "")[:31]
    for bad, good in _BAD.items():
        text = text.replace(bad, good)
    return text


def excel_rows(wb, sheets, session_id):
    # Case-insensitive, like Excel itself: session "Testing" may live on tab
    # "testing". A case-sensitive lookup reports a healthy session as having no
    # sheet — that false positive is exactly what this sweep first reported.
    actual = S.resolve_sheet_title(sheets, sanitize_sheet_name(session_id))
    if actual is None:
        return None
    rows = []
    for r in wb[actual].iter_rows(min_row=2, values_only=True):
        if not r or not any(r):
            continue
        c = (list(r) + [None] * 9)[:9]
        rows.append({"name": c[0], "mobile": c[1], "session": c[4], "batch": c[5],
                     "rounds": c[6], "remark": c[7], "date": c[8]})
    return rows


def main():
    conn = sqlite3.connect(DB)
    conn.row_factory = sqlite3.Row
    wb = load_workbook(LIVE, read_only=True, data_only=True)
    sheets = set(wb.sheetnames)

    totals = {S.MISSING_IN_EXCEL: 0, S.MISSING_IN_DB: 0,
              S.FIELD_MISMATCH: 0, S.DUPLICATE_IN_EXCEL: 0}
    no_sheet, repairable, dirty, checked = [], [], 0, 0

    for s in conn.execute("SELECT session_id, centre_id FROM attendance "
                          "GROUP BY session_id, centre_id"):
        sid, cid = s["session_id"], s["centre_id"]
        checked += 1
        db_rows = [dict(r) for r in conn.execute(
            "SELECT a.mobile, a.date, a.rounds, a.remark, u.name FROM attendance a "
            "LEFT JOIN users u ON u.mobile = a.mobile "
            "WHERE a.session_id = ? AND a.centre_id = ?", (sid, cid))]
        ex = excel_rows(wb, sheets, sid)
        rep = S.audit(db_rows, ex or [], sheet_exists=ex is not None)
        for k in totals:
            totals[k] += len(rep[k])
        if ex is None:
            no_sheet.append((cid, sid, len(db_rows)))
        elif rep[S.MISSING_IN_EXCEL]:
            repairable.append((cid, sid, len(rep[S.MISSING_IN_EXCEL])))
        if not S.is_clean(rep):
            dirty += 1

    print(f"sessions checked: {checked}   not in sync: {dirty}")
    print("\ndrift totals")
    for k, v in totals.items():
        print(f"  {k:20} {v}")
    print(f"\nsessions with NO sheet at all: {len(no_sheet)}")
    for cid, sid, n in no_sheet:
        print(f"   {cid:20} {sid!r}  ({n} db rows)")
    print(f"\nsessions whose sheet exists but is missing rows: {len(repairable)}")
    for cid, sid, n in repairable[:20]:
        print(f"   {cid:20} {sid!r}  ({n} missing)")
    conn.close()


if __name__ == "__main__":
    main()
