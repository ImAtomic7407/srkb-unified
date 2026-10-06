"""Behaviour tests for excel_lock_guard.

Covers, in order of how much they matter:
  T5  a holder with no /automation flag is never killed and its data survives  <- the safety property
  T4  an automation-owned holder is killed, the ~$ swept, and the write retried and lands
  T3  safe_excel_write self-heals a blocked save with no caller change
  T2  ~$ files: stale ones are removed, live ones are left alone
  T1  the probe itself, and the refusal when no Excel process explains the lock

Only the *identity* of the holder is simulated: excel_processes() is patched to
report a real, really-holding, really-killable process as EXCEL.EXE. Everything
else -- the exclusive OS handle, the PermissionError, the kill, the release wait,
the sweep, the retry -- is the real thing.
"""
import os
import shutil
import subprocess
import sys
import time

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, "..")))

import excel_lock_guard as G  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

RIG = os.path.join(HERE, "rig.xlsx")
LOG = []
failures = []


def log(msg):
    LOG.append(msg)


def check(name, cond, detail=""):
    print(("  PASS  " if cond else "  FAIL  ") + name + (" :: " + detail if detail else ""))
    if not cond:
        failures.append(name)


def fresh_rig():
    for f in os.listdir(HERE):
        if f.startswith("~$"):
            try:
                os.remove(os.path.join(HERE, f))
            except OSError:
                pass
    wb = Workbook()
    ws = wb.active
    ws.title = "MASTER_ALL_SESSIONS"
    ws.append(["Name", "Number"])
    ws.append(["Seed", "9999999999"])
    wb.save(RIG)


def start_holder(path, seconds=90):
    proc = subprocess.Popen([sys.executable, os.path.join(HERE, "holder.py"), path, str(seconds)],
                            stdout=subprocess.PIPE, text=True)
    line = proc.stdout.readline().strip()
    assert line == "HOLDING", "holder failed to take the lock: %r" % line
    return proc


def patch_processes(pid, is_automation, has_window=False):
    G.excel_processes = lambda timeout=20: (
        [{"pid": pid, "has_window": has_window, "is_automation": is_automation}]
        if _alive(pid) else [])


def _alive(pid):
    out = subprocess.run(["powershell", "-NoProfile", "-Command",
                          "if (Get-Process -Id %d -ErrorAction SilentlyContinue) {'Y'} else {'N'}" % pid],
                         capture_output=True, text=True)
    return out.stdout.strip() == "Y"


_REAL_PROCESSES = G.excel_processes


# ── T1: probe, and refusal when nothing explains the lock ────────────────────
print("\nT1  probe + unknown holder")
fresh_rig()
check("writable file probes clean", not G.is_locked(RIG))
holder = start_holder(RIG)
try:
    check("held file probes locked", G.is_locked(RIG))
    check("lock_error is a PermissionError-class error", G._is_lock_error(G.lock_error(RIG)),
          type(G.lock_error(RIG)).__name__)
    G.excel_processes = lambda timeout=20: []   # no Excel at all
    status = G.clear_excel_lock(RIG, log=log)
    check("refuses when no Excel explains the lock", status == G.CLEAR_UNKNOWN_HOLDER, status)
    check("holder was NOT killed", _alive(holder.pid))
finally:
    holder.kill(); holder.wait()
    G.excel_processes = _REAL_PROCESSES
time.sleep(0.4)
check("probes clean again once released", not G.is_locked(RIG))


# ── T2: ~$ sweep ─────────────────────────────────────────────────────────────
print("\nT2  ~$ lock-file sweep")
fresh_rig()
stale = os.path.join(HERE, "~$rig.xlsx")
open(stale, "wb").write(b"stale")
removed = G.sweep_stale_lock_files(HERE, log)
check("stale ~$ removed", not os.path.exists(stale), str(removed))

live = os.path.join(HERE, "~$live.xlsx")
open(live, "wb").write(b"live")
live_holder = start_holder(live)
try:
    removed = G.sweep_stale_lock_files(HERE, log)
    check("live ~$ left alone", os.path.exists(live), str(removed))
finally:
    live_holder.kill(); live_holder.wait()
    time.sleep(0.4)
    G.sweep_stale_lock_files(HERE, log)


# ── T3/T4: automation-owned holder -> killed, swept, write lands ─────────────
print("\nT4  automation-owned holder is cleared and the write completes")
fresh_rig()
open(os.path.join(HERE, "~$rig.xlsx"), "wb").write(b"stale")
holder = start_holder(RIG)
patch_processes(holder.pid, is_automation=True)
journal_rows = []
try:
    check("write is genuinely blocked first", G.is_locked(RIG))
    wb = load_workbook(RIG)
    wb["MASTER_ALL_SESSIONS"].append(["Recovered", "8888888888"])
    G.safe_excel_write(
        RIG, lambda: wb.save(RIG), log=log,
        journal=lambda **kw: journal_rows.append(kw), reason="test_attendance")
    check("save succeeded after self-heal", True)
    check("holder was killed", not _alive(holder.pid))
    check("stale ~$ swept", not os.path.exists(os.path.join(HERE, "~$rig.xlsx")))
    check("journal recorded self_healed",
          any(r.get("status") == "self_healed" for r in journal_rows), str(journal_rows))
    rows = [tuple(r) for r in load_workbook(RIG).active.iter_rows(values_only=True)]
    check("the row actually landed on disk", ("Recovered", "8888888888") in rows, str(rows))
except Exception as e:
    check("save succeeded after self-heal", False, repr(e))
finally:
    if holder.poll() is None:
        holder.kill(); holder.wait()
    G.excel_processes = _REAL_PROCESSES


# ── T5: THE SAFETY PROPERTY -- a non-automation holder is never killed ───────
print("\nT5  human-owned holder (no /automation) is refused, unsaved work survives")
fresh_rig()
holder = start_holder(RIG)
patch_processes(holder.pid, is_automation=False, has_window=True)
journal_rows = []
try:
    status = G.clear_excel_lock(RIG, log=log,
                                journal=lambda **kw: journal_rows.append(kw), reason="test")
    check("refuses to clear", status == G.CLEAR_HELD_BY_HUMAN, status)
    check("the human's Excel is STILL RUNNING", _alive(holder.pid))
    check("journal recorded needs_attention",
          any(r.get("status") == "needs_attention" for r in journal_rows), str(journal_rows))

    wb = load_workbook(RIG)
    wb["MASTER_ALL_SESSIONS"].append(["ShouldNotLand", "7777777777"])
    raised = None
    try:
        G.safe_excel_write(RIG, lambda: wb.save(RIG), log=log, reason="test")
    except OSError as e:
        raised = e
    check("safe_excel_write re-raises so the caller queues the row", raised is not None,
          type(raised).__name__ if raised else "no exception")
    check("the human's Excel survived the write attempt", _alive(holder.pid))
finally:
    if holder.poll() is None:
        holder.kill(); holder.wait()
    G.excel_processes = _REAL_PROCESSES
time.sleep(0.4)
rows = [tuple(r) for r in load_workbook(RIG).active.iter_rows(values_only=True)]
check("workbook on disk was not corrupted by the refused write",
      ("Seed", "9999999999") in rows, str(rows))


# ── T7: THE REGRESSION -- never kill a COM write of our own that is in flight ─
# An orphan and a healthy COM write are indistinguishable from outside: both are
# an /automation Excel holding the workbook. The COM subprocess runs OUTSIDE
# EXCEL_WRITE_LOCK, so a second thread can hit a blocked save while the
# attendance writer is mid-$wb.Save(). Killing then would destroy in-flight work
# -- strictly worse than the old behaviour, where that save just failed into the
# pending store. The in-flight counter is what separates the two cases.
print("\nT7  a COM write of ours that is in flight is never killed")
fresh_rig()
holder = start_holder(RIG)
patch_processes(holder.pid, is_automation=True)   # looks exactly like an orphan
journal_rows = []
try:
    with G.com_write():                            # ...but a COM write is running
        check("counter reports the write in flight", G.com_writes_in_flight() == 1,
              str(G.com_writes_in_flight()))
        status = G.clear_excel_lock(RIG, log=log,
                                    journal=lambda **kw: journal_rows.append(kw), reason="test")
        check("refuses while our own COM write is running", status == G.CLEAR_BUSY, status)
        check("the in-flight writer was NOT killed", _alive(holder.pid))
        check("journal recorded skipped, not self_healed",
              any(r.get("status") == "skipped" for r in journal_rows)
              and not any(r.get("status") == "self_healed" for r in journal_rows),
              str(journal_rows))
    check("counter clears when the write finishes", G.com_writes_in_flight() == 0)
    # Once the write is done, the same holder IS reapable.
    status = G.clear_excel_lock(RIG, log=log, reason="test")
    check("same holder is reaped once no write is in flight", status == G.CLEAR_SELF_HEALED, status)
finally:
    if holder.poll() is None:
        holder.kill(); holder.wait()
    G.excel_processes = _REAL_PROCESSES

# The counter must survive an exception so a timeout handler can still reap.
try:
    with G.com_write():
        raise RuntimeError("simulated COM timeout")
except RuntimeError:
    pass
check("counter is decremented even when the COM write raises", G.com_writes_in_flight() == 0,
      str(G.com_writes_in_flight()))


# ── T6: the discriminator against a REAL EXCEL.EXE ───────────────────────────
# T1-T5 stand in for the holder process; this one checks the single fact
# everything else rests on -- that excel_processes() correctly labels a genuine
# COM-created Excel as an automation instance. Skipped if Excel is unavailable.
print("\nT6  real EXCEL.EXE is identified as an automation instance")
spawn = subprocess.Popen(
    ["powershell", "-NoProfile", "-Command",
     "$x = New-Object -ComObject Excel.Application; $x.Visible = $false; "
     "Write-Output READY; Start-Sleep -Seconds 25"],
    stdout=subprocess.PIPE, text=True)
try:
    ready = spawn.stdout.readline().strip()
    if ready != "READY":
        print("  SKIP  Excel COM unavailable in this environment")
    else:
        time.sleep(2.0)
        procs = _REAL_PROCESSES()
        check("a real EXCEL.EXE was seen", len(procs) > 0, str(procs))
        check("it is labelled is_automation", any(p.get("is_automation") for p in procs),
              str(procs))
        check("its command line carries /automation -Embedding",
              all(p.get("is_automation") for p in procs), str(procs))
finally:
    spawn.kill(); spawn.wait()
    subprocess.run(["powershell", "-NoProfile", "-Command",
                    "Get-Process -Name EXCEL -ErrorAction SilentlyContinue | "
                    "Where-Object { $_.MainWindowHandle -eq 0 } | Stop-Process -Force"],
                   capture_output=True)


# Leave nothing behind — this runs against a real directory on a live machine.
for _f in os.listdir(HERE):
    if _f == "rig.xlsx" or _f.startswith("~$"):
        try:
            os.remove(os.path.join(HERE, _f))
        except OSError:
            pass

print("\n" + ("ALL PASS" if not failures else "FAILURES: " + ", ".join(failures)))
sys.exit(1 if failures else 0)
